"""Thin helper layer over python-docx: letterhead header, draft banner, page numbers, tables, figures, captions."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PLACEHOLDER_RE = re.compile(r"(\[(?:P\.G\.|REVIEWER|P\.E\.)[^\]]*\])")
BODY_FONT = "Times New Roman"


def _set_cell_shading(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _add_field(run, instr: str):
    for tag, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
            run._r.append(el)
        elif text == instr:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = f" {instr} "
            run._r.append(it)
        else:
            t = OxmlElement("w:t")
            t.text = text
            run._r.append(t)


class DocBuilder:
    def __init__(self, letterhead: dict, banner: str | None = None, footer_text: str = ""):
        self.doc = Document()
        st = self.doc.styles["Normal"]
        st.font.name = BODY_FONT
        st.font.size = Pt(11)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        for lvl, size in ((1, 13), (2, 12), (3, 11)):
            h = self.doc.styles[f"Heading {lvl}"]
            h.font.name = BODY_FONT
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = RGBColor(0, 0, 0)
            h.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        sec = self.doc.sections[0]
        sec.left_margin = sec.right_margin = Inches(1.0)
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(0.9)
        self.letterhead = letterhead
        self.banner = banner
        self.footer_text = footer_text
        self._decorate_section(sec)
        self.fig_no = 0
        self.tab_no = 0

    # ---------- header / footer ----------
    def _decorate_section(self, sec):
        hdr = sec.header
        hdr.is_linked_to_previous = False
        p = hdr.paragraphs[0]
        p.text = ""
        r = p.add_run(self.letterhead.get("firm", ""))
        r.bold = True
        r.font.size = Pt(11)
        lines = list(self.letterhead.get("lines", []))
        if self.letterhead.get("registration_line"):
            lines.append(self.letterhead["registration_line"])
        if lines:
            p2 = hdr.add_paragraph()
            rr = p2.add_run("  |  ".join(lines))
            rr.font.size = Pt(7.5)
            p2.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        if self.banner:
            pb = hdr.add_paragraph()
            pb.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rb = pb.add_run(self.banner)
            rb.bold = True
            rb.font.size = Pt(10)
            rb.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        ftr = sec.footer
        ftr.is_linked_to_previous = False
        pf = ftr.paragraphs[0]
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.text = ""
        rf = pf.add_run((self.footer_text + "    Page ") if self.footer_text else "Page ")
        rf.font.size = Pt(8)
        r1 = pf.add_run()
        r1.font.size = Pt(8)
        _add_field(r1, "PAGE")
        r2 = pf.add_run(" of ")
        r2.font.size = Pt(8)
        r3 = pf.add_run()
        r3.font.size = Pt(8)
        _add_field(r3, "NUMPAGES")

    # ---------- text ----------
    def heading(self, text, level=1):
        return self.doc.add_heading(text, level=level)

    def para(self, text, bold=False, italic=False, size=None, align=None, space_after=6, style=None):
        p = self.doc.add_paragraph(style=style) if style else self.doc.add_paragraph()
        self._runs(p, text, bold=bold, italic=italic, size=size)
        p.paragraph_format.space_after = Pt(space_after)
        if align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "justify":
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return p

    def _runs(self, p, text, bold=False, italic=False, size=None):
        for part in PLACEHOLDER_RE.split(text):
            if not part:
                continue
            r = p.add_run(part)
            r.bold = bold
            r.italic = italic
            if size:
                r.font.size = Pt(size)
            if PLACEHOLDER_RE.fullmatch(part):
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                r.bold = True

    def paragraphs(self, text, align="justify"):
        """Render a block of text: blank-line separated paragraphs; lines starting with '- ' become bullets;
        lines starting with 'N. ' become numbered items."""
        for block in re.split(r"\n\s*\n", text.strip()):
            lines = block.strip().split("\n")
            if all(ln.strip().startswith("- ") for ln in lines):
                for ln in lines:
                    self.bullet(ln.strip()[2:])
            elif all(re.match(r"^\d+\.\s", ln.strip()) for ln in lines):
                for ln in lines:
                    self.numbered(re.sub(r"^\d+\.\s", "", ln.strip()))
            else:
                self.para(" ".join(ln.strip() for ln in lines), align=align)

    def bullet(self, text):
        p = self.doc.add_paragraph(style="List Bullet")
        self._runs(p, text)
        p.paragraph_format.space_after = Pt(2)
        return p

    def numbered(self, text):
        p = self.doc.add_paragraph(style="List Number")
        self._runs(p, text)
        p.paragraph_format.space_after = Pt(2)
        return p

    def page_break(self):
        self.doc.add_page_break()

    # ---------- tables / figures ----------
    def next_table(self):
        self.tab_no += 1
        return self.tab_no

    def next_figure(self):
        self.fig_no += 1
        return self.fig_no

    def table(self, number, caption, header, rows, col_widths_in=None, font_pt=8, note=None, header_fill="D9E2F3",
              red_cells=None):
        cap = self.doc.add_paragraph()
        r = cap.add_run(f"Table {number}: {caption}")
        r.bold = True
        r.font.size = Pt(10)
        cap.paragraph_format.keep_with_next = True
        cap.paragraph_format.space_after = Pt(2)
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]
            c.text = ""
            rr = c.paragraphs[0].add_run(str(h))
            rr.bold = True
            rr.font.size = Pt(font_pt)
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            _set_cell_shading(c, header_fill)
        red_cells = red_cells or set()
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for ci, val in enumerate(row):
                cells[ci].text = ""
                rr = cells[ci].paragraphs[0].add_run("" if val is None else str(val))
                rr.font.size = Pt(font_pt)
                if (ri, ci) in red_cells:
                    rr.font.color.rgb = RGBColor(0xC0, 0, 0)
                    rr.bold = True
                cells[ci].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if ci > 0 else WD_ALIGN_PARAGRAPH.LEFT
        if col_widths_in:
            t.autofit = False
            for row in t.rows:
                for ci, w in enumerate(col_widths_in):
                    if ci < len(row.cells):
                        row.cells[ci].width = Inches(w)
        if note:
            pn = self.doc.add_paragraph()
            rn = pn.add_run(note)
            rn.italic = True
            rn.font.size = Pt(8)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def figure(self, number, caption, path, width_in=6.0):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width_in))
        p.paragraph_format.keep_with_next = True
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(f"Figure {number}: {caption}")
        r.bold = True
        r.font.size = Pt(9.5)
        cap.paragraph_format.space_after = Pt(10)

    def equation(self, path, label):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), height=Inches(0.55))
        p.add_run(f"        ({label})").font.size = Pt(9)

    # ---------- sections ----------
    def landscape(self):
        sec = self.doc.add_section()
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
        sec.left_margin = sec.right_margin = Inches(0.8)
        self._decorate_section(sec)
        return sec

    def portrait(self):
        sec = self.doc.add_section()
        sec.orientation = WD_ORIENT.PORTRAIT
        if sec.page_width > sec.page_height:
            sec.page_width, sec.page_height = sec.page_height, sec.page_width
        sec.left_margin = sec.right_margin = Inches(1.0)
        self._decorate_section(sec)
        return sec

    def save(self, path: Path):
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))
        return str(path)
