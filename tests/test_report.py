import shutil
from pathlib import Path

import pytest
from docx import Document

from hydrostudy.pipeline import Project
from hydrostudy.report.lint import lint_sections
from tests.conftest import ROOT


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    src = ROOT / "examples" / "black_oak_well_2"
    dst = tmp_path_factory.mktemp("bo") / "black_oak_well_2"
    shutil.copytree(src, dst, ignore=shutil.ignore_patterns("build"))
    p = Project(dst)
    p.run_analysis()
    p.run_figures()
    p.artifacts["checklist"] = p.run_checklist()
    from hydrostudy.report.assemble import build_report
    result = build_report(p, pdf=False)
    return p, result


def test_docx_structure(built):
    _, result = built
    assert result["lint"]["ok"]
    doc = Document(result["docx"])
    heads = [par.text for par in doc.paragraphs if par.style.name.startswith("Heading 1")]
    expected = ["1. Well Spacing", "2. Well Construction Details", "3. General Hydrogeology", "4. Site-Specific Hydrogeology",
                "5. Water Quality", "6. Interference Analysis", "7. Summary and Professional Opinion", "References"]
    idx = [heads.index(h) for h in expected]
    assert idx == sorted(idx)
    text = "\n".join(par.text for par in doc.paragraphs)
    assert "{{" not in text and "}}" not in text
    import re
    caps = [par.text for par in doc.paragraphs if re.match(r"^Figure \d+: ", par.text)]
    nums = [int(c.split(":")[0].split()[1]) for c in caps]
    assert nums == list(range(1, len(nums) + 1))
    assert "DRAFT - NOT SEALED" in doc.sections[0].header.paragraphs[-1].text
    assert len(result["placeholders"]) == 4
    assert "98.7" in text and "121.5" in text and "52.31" in text


def test_spacing_states_the_rule_and_discloses_its_basis(built):
    """Lone Star spacing is operator-attested: the report may state the rule, but must say what that rests on."""
    _, result = built
    text = "\n".join(par.text for par in Document(result["docx"]).paragraphs)
    assert "complies with the spacing rule" in text
    assert "[REVIEWER TO CONFIRM: the current District spacing multiplier" not in text
    # The sealing professional must be able to find the basis without opening the repo.
    assert "stated on the attestation of" in text
    assert "not on a reading of the District Rules" in text


def test_checklist(built):
    p, _ = built
    cl = p.artifacts["checklist"]
    assert cl["counts"]["satisfied"] >= 17
    assert all(i["status"] in ("satisfied", "needs_professional_input") for i in cl["items"])


def test_lint_catches_unknown_number():
    ctx = {"a": "drawdown of 98.7 ft"}
    assert lint_sections({"s": "drawdown of 98.7 ft"}, ctx)["ok"]
    bad = lint_sections({"s": "drawdown of 99.9 ft"}, ctx)
    assert not bad["ok"] and bad["problems"][0]["number"] == "99.9"


def test_pdf_conversion(built):
    from hydrostudy.report.pdf import convert_to_pdf, soffice_available
    if not soffice_available():
        pytest.skip("LibreOffice Writer not available")
    _, result = built
    pdf = convert_to_pdf(Path(result["docx"]))
    assert pdf is not None and pdf.stat().st_size > 50_000
